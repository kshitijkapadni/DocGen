from docx import Document
import os
import uuid


from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def create_structured_doc(content, version: str, output_path: str):
    doc = Document()

    # Title
    doc.add_heading("📄 Project Documentation", 0)
    doc.add_paragraph(f"Version: {version}")

    # Table of Contents placeholder
    doc.add_page_break()
    doc.add_paragraph("Table of Contents").style = 'Heading 1'

    # Content blocks
    for i, section in enumerate(content):
        doc.add_page_break()
        doc.add_heading(f"Section {i+1}", level=1)
        doc.add_paragraph(section, style='Normal')

    # Save it
    doc.save(output_path)

def generate_docx(payload: dict) -> str:
    doc = Document()
    doc.add_heading("Application Documentation", 0)
    
    description = payload.get("description", "No description provided.")
    doc.add_paragraph(f"App Description: {description}")
    
    for entry in payload.get("chat", []):
        doc.add_paragraph(f"{entry['sender']}: {entry['message']}")
    
    path = f"/tmp/generated_{uuid.uuid4().hex}.docx"
    doc.save(path)
    return path
